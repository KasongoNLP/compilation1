import os
from copy import copy
from openpyxl import load_workbook, Workbook
import customtkinter

DOSSIER = "fichiers"
FICHIER_SORTIE = "checklist.xlsx"



compilation = Workbook()
class interface(ctk.CTk) :
    def __init__(self) :
        self.title(APP_NAME)
        self.geometry("1200x760")
        self.minsize(1050, 680)







# compilation.remove(compilation.active)

i = interface()
for fichier in os.listdir(DOSSIER):

    if not fichier.endswith((".Xlsx", ".xlsm")):
        continue
    if fichier == FICHIER_SORTIE:
        continue

    chemin = os.path.join(DOSSIER, fichier)
    workbook_source = load_workbook(chemin)
    feuille_source = workbook_source.worksheets[0]
    nom_sheet = os.path.splitext(fichier)[0]
    nom_sheet = nom_sheet[:31]


    feuille_destination = compilation.create_sheet(nom_sheet)


    for ligne in feuille_source.iter_rows():
        for cellule in ligne:
            nouvelle_cellule = feuille_destination[cellule.coordinate]
            nouvelle_cellule.value = cellule.value


            if cellule.has_style:
                nouvelle_cellule.font = copy(cellule.font)
                nouvelle_cellule.fill = copy(cellule.fill)
                nouvelle_cellule.border = copy(cellule.border)
                nouvelle_cellule.alignment = copy(cellule.alignment)
                nouvelle_cellule.number_format = cellule.number_format
                nouvelle_cellule.protection = copy(cellule.protection)


    for colonne, dimension in feuille_source.column_dimensions.items():
        feuille_destination.column_dimensions[colonne].width = dimension.width


    for ligne, dimension in feuille_source.row_dimensions.items():
        feuille_destination.row_dimensions[ligne].height = dimension.height


    for plage in feuille_source.merged_cells.ranges:
        feuille_destination.merge_cells(str(plage))

    print(f" {fichier} copié")


chemin_final = os.path.join(DOSSIER, FICHIER_SORTIE)
compilation.save(chemin_final)

print("Compilation terminee !")
print(f"Fichier cree : {chemin_final}")