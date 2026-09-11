import customtkinter as ctk
from PIL import Image
from copy import copy
from openpyxl import load_workbook, Workbook
import customtkinter
import os
from tkinter import messagebox
from datetime import date
# ============================================================
# CONFIGURATION De l'affichage
# ============================================================

ctk.set_appearance_mode("light")
ctk.set_default_color_theme("blue")

# ============================================================
# FENÊTRE PRINCIPALE
# ============================================================

fenetre = ctk.CTk()

fenetre.title("Dailly Checking")
fenetre.geometry("800x500")
fenetre.resizable(True, True)

# Icône de la fenêtre
fenetre.iconbitmap("icone.ico")

# ============================================================
# FONCTION CONNEXION
# ============================================================

def ouvrir_connexion():
    # ========================================================
    # FENÊTRE POPUP
    # ========================================================
    modal = ctk.CTkToplevel(fenetre)
    modal.title("Connexion")
    modal.geometry("400x300")
    modal.resizable(False, False)

    # Popup au-dessus de la fenêtre principale
    modal.transient(fenetre)

    # Bloquer la fenêtre principale
    modal.grab_set()



    # ========================================================
    # TITRE
    # ========================================================
    titre = ctk.CTkLabel(modal,text="Connexion",font=ctk.CTkFont(size=24,weight="bold"))
    titre.pack(pady=(25, 15))


    # ========================================================
    # NOM UTILISATEUR
    # ========================================================
    utilisateur = ctk.CTkEntry(modal,width=260,height=35,placeholder_text="Nom d'utilisateur")
    utilisateur.pack(pady=5)


    # ========================================================
    # MOT DE PASSE
    # ========================================================

    mot_de_passe = ctk.CTkEntry(modal,width=260,height=35,placeholder_text="Mot de passe",show="*")
    mot_de_passe.pack(pady=5)


    # ========================================================
    # MESSAGE D'ERREUR
    # ========================================================
    erreur = ctk.CTkLabel(modal,text="",text_color="red")
    erreur.pack(pady=3)


    # ========================================================
    # FONCTION DE CONNEXION
    # ========================================================

    def connecter():

        user = utilisateur.get()
        password = mot_de_passe.get()

        if user == "steve" and password == "steve":

            print("Connexion réussie")
            afficher_accueil()

            modal.grab_release()
            modal.destroy()

        else:

            erreur.configure(
                text="Identifiants incorrects"
            )





    # ========================================================
    # BOUTON
    # ========================================================
    bouton = ctk.CTkButton(modal,text="Se connecter",width=150,height=35,fg_color="red",command=connecter)
    bouton.pack(pady=8)







# ================================================================================
# OUVRIR LE POPUP AUTOMATIQUEMENT
# ================================================================================
fenetre.after(100,ouvrir_connexion)
# ================================================================================


#........... Gestion de frame .................................................
conteneur = ctk.CTkFrame(fenetre,fg_color="#F3F4F6",corner_radius=20)
conteneur.place(x=0,y=0,relheight=1)

#..............................................................................
conteneur1 = ctk.CTkFrame(fenetre,fg_color="#F8FAFC",corner_radius=18,border_width=1,border_color="#E5E7EB")
conteneur1.place(x=200,y=10,relwidth=0.82,relheight=0.96)

def afficher_accueil():

    # Nettoyer le conteneur principal
    for widget in conteneur.winfo_children():
        widget.destroy()

    # ========================================================
    # STRUCTURE PRINCIPALE
    # ========================================================

    frame_accueil = ctk.CTkFrame(
        conteneur,
        fg_color="transparent"
    )

    frame_accueil.pack(
        fill="both",
        expand=True
    )

    # ========================================================
    # MENU GAUCHE
    # ========================================================

    menu = ctk.CTkFrame(
        frame_accueil,
        width=190,
        corner_radius=18,
        fg_color="white"
    )

    menu.place(
        x=0,
        y=0,
        relheight=1
    )
    menu.pack_propagate(False)

    # ========================================================
    # TITRE
    # ========================================================

    titre_menu = ctk.CTkLabel(
        menu,
        text="DAILLY\nCHECKING",
        text_color="#1F2937",
        font=ctk.CTkFont(
            size=19,
            weight="bold"
        )
    )

    titre_menu.pack(
        pady=(22, 12)
    )

    # ========================================================
    # PHOTO DE PROFIL
    # ========================================================

    try:

        image_profil = Image.open("profil.png")
        image_profil = image_profil.resize((70, 70))

        photo_profil = ctk.CTkImage(
            light_image=image_profil,
            dark_image=image_profil,
            size=(70, 70)
        )

        label_photo = ctk.CTkLabel(
            menu,
            text="",
            image=photo_profil
        )

    except Exception:

        label_photo = ctk.CTkLabel(
            menu,
            text="S",
            width=70,
            height=70,
            corner_radius=35,
            fg_color="#E5E7EB",
            text_color="#374151",
            font=ctk.CTkFont(
                size=28,
                weight="bold"
            )
        )

    label_photo.pack(
        pady=(5, 5)
    )

    # ========================================================
    # NOM UTILISATEUR
    # ========================================================

    nom_utilisateur = ctk.CTkLabel(
        menu,
        text="Steve",
        text_color="#1F2937",
        font=ctk.CTkFont(
            size=14,
            weight="bold"
        )
    )

    nom_utilisateur.pack(
        pady=(2, 18)
    )

    # ========================================================
    # SÉPARATION
    # ========================================================

    separation = ctk.CTkFrame(
        menu,
        height=1,
        fg_color="#E5E7EB"
    )

    separation.pack(
        fill="x",
        padx=20,
        pady=(0, 15)
    )

    # ========================================================
    # BOUTONS
    # ========================================================

    bouton_profil = ctk.CTkButton(
        menu,
        text="Profil",
        height=38,
        corner_radius=10,
        anchor="w",
        fg_color="transparent",
        hover_color="#F1F5F9",
        text_color="#374151"
    )

    bouton_profil.pack(
        fill="x",
        padx=12,
        pady=3
    )

    bouton_compilation = ctk.CTkButton(
        menu,
        text="Compilation",
        height=38,
        corner_radius=10,
        anchor="w",
        fg_color="#EFF6FF",
        hover_color="#DBEAFE",
        text_color="#2563EB",
        command=afficher_compilation
    )

    bouton_compilation.pack(
        fill="x",
        padx=12,
        pady=3
    )

    bouton_historyque = ctk.CTkButton(
        menu,
        text="Historique",
        height=38,
        corner_radius=10,
        anchor="w",
        fg_color="transparent",
        hover_color="#F1F5F9",
        text_color="#374151",
        command=afficher_historyque
    )

    bouton_historyque.pack(
        fill="x",
        padx=12,
        pady=3
    )

    bouton_setting = ctk.CTkButton(
        menu,
        text="Setting",
        height=38,
        corner_radius=10,
        anchor="w",
        fg_color="transparent",
        hover_color="#F1F5F9",
        text_color="#374151"
    )

    bouton_setting.pack(
        fill="x",
        padx=12,
        pady=3
    )

    bouton_email = ctk.CTkButton(
        menu,
        text="Email",
        height=38,
        corner_radius=10,
        anchor="w",
        fg_color="transparent",
        hover_color="#F1F5F9",
        text_color="#374151"
    )

    bouton_email.pack(
        fill="x",
        padx=12,
        pady=3
    )

    # ========================================================
    # ESPACE FLEXIBLE
    # ========================================================

    espace = ctk.CTkFrame(
        menu,
        fg_color="transparent"
    )

    espace.pack(
        fill="both",
        expand=True
    )

    # ========================================================
    # DÉCONNEXION
    # ========================================================

    bouton_deconnexion = ctk.CTkButton(
        menu,
        text="Déconnexion",
        height=38,
        corner_radius=10,
        anchor="w",
        fg_color="#FEF2F2",
        hover_color="#FEE2E2",
        text_color="#DC2626"
    )

    bouton_deconnexion.pack(
        fill="x",
        padx=5,
        pady=(3, 18)
    )



#Gestion de fenêtre affichage de compilation 
def afficher_compilation():


    # ========================================================
    # NETTOYER LA ZONE DE CONTENU
    # ========================================================
    for widget in conteneur1.winfo_children():
        widget.destroy()


    # ========================================================
    # VARIABLES
    # ========================================================
    fichiers_selectionnes = []


    # ========================================================
    # CONTENEUR PRINCIPAL DE LA PAGE
    # ========================================================
    page = ctk.CTkFrame(conteneur1,fg_color="transparent")

    page.pack(fill="both",expand=True,padx=20,pady=20)


    # ========================================================
    # EN-TÊTE
    # ========================================================
    header = ctk.CTkFrame(page,fg_color="transparent")
    header.pack(fill="x",pady=(0, 15))


    titre = ctk.CTkLabel(
        header,
        text="Compilation",
        text_color="#1F2937",
        font=ctk.CTkFont(
            size=26,
            weight="bold"
        )
    )

    titre.pack(
        anchor="w"
    )


    description = ctk.CTkLabel(
        header,
        text="Sélectionnez les fichiers Excel à compiler",
        text_color="#6B7280",
        font=ctk.CTkFont(
            size=13
        )
    )

    description.pack(
        anchor="w",
        pady=(3, 0)
    )


    # ========================================================
    # BARRE D'ACTIONS
    # ========================================================
    actions = ctk.CTkFrame(page,fg_color="white",corner_radius=15)
    actions.pack(fill="x",pady=(0, 15))


    # ========================================================
    # AJOUTER DES FICHIERS
    # ========================================================

    def ajouter_fichiers():
        fichiers = ctk.filedialog.askopenfilenames(
            title="Sélectionner les fichiers Excel",
            filetypes=[
                ("Fichiers Excel", "*.xlsx *.xlsm"),
                ("Fichiers XLSX", "*.xlsx"),
                ("Fichiers XLSM", "*.xlsm")
            ]
        )

        if not fichiers:
            return

        for fichier in fichiers:
            if fichier not in fichiers_selectionnes:
                fichiers_selectionnes.append(fichier)

        afficher_liste_fichiers()


    bouton_ajouter = ctk.CTkButton(
        actions,
        text="+  Ajouter des fichiers",
        width=180,
        height=40,
        corner_radius=10,
        fg_color="#2563EB",
        hover_color="#1D4ED8",
        command=ajouter_fichiers
    )

    bouton_ajouter.pack(
        side="left",
        padx=15,
        pady=15
    )


    # ========================================================
    # RECHERCHER UNE COMPILATION
    # ========================================================

    def rechercher_compilation():

        fichier = ctk.filedialog.askopenfilename(
            title="Rechercher une compilation",
            filetypes=[
                ("Fichier Excel", "*.xlsx"),
                ("Fichier XLSM", "*.xlsm")
            ]
        )

        if not fichier:
            return

        # Ouvrir le fichier avec l'application Excel par défaut
        try:
            os.startfile(fichier)
        except Exception as e:
            print(f"Impossible d'ouvrir le fichier : {e}")


    bouton_rechercher = ctk.CTkButton(actions,text="Ouvrir une compilation",width=90,height=10,corner_radius=10,
        fg_color="#E5E7EB",
        hover_color="#D1D5DB",
        text_color="#374151",
        command=rechercher_compilation
    )

    bouton_rechercher.pack(
        side="left",
        padx=5,
        pady=15
    )


    # ========================================================
    # ESPACE DES FICHIERS
    # ========================================================

    carte_fichiers = ctk.CTkFrame(
        page,
        fg_color="white",
        corner_radius=15
    )

    carte_fichiers.pack(
        fill="both",
        expand=True
    )


    # ========================================================
    # EN-TÊTE DE LA LISTE
    # ========================================================

    header_fichiers = ctk.CTkFrame(
        carte_fichiers,
        fg_color="transparent"
    )

    header_fichiers.pack(
        fill="x",
        padx=20,
        pady=(18, 10)
    )


    titre_fichiers = ctk.CTkLabel(
        header_fichiers,
        text="Fichiers à compiler",
        text_color="#1F2937",
        font=ctk.CTkFont(
            size=17,
            weight="bold"
        )
    )

    titre_fichiers.pack(
        side="left"
    )


    nombre_fichiers = ctk.CTkLabel(
        header_fichiers,
        text="0 fichier",
        text_color="#6B7280",
        font=ctk.CTkFont(
            size=12
        )
    )

    nombre_fichiers.pack(
        side="right"
    )


    # ========================================================
    # ZONE SCROLLABLE
    # ========================================================

    liste = ctk.CTkScrollableFrame(
        carte_fichiers,
        fg_color="#F8FAFC",
        corner_radius=10
    )

    liste.pack(
        fill="both",
        expand=True,
        padx=20,
        pady=(0, 15)
    )



    label_type = ctk.CTkLabel(
        page,
        text="Type de compilation",
        font=("Arial", 14, "bold"),
        text_color="#111827"
    )
    label_type.pack(anchor="w", padx=30, pady=(20, 5))

    select_type = ctk.CTkComboBox(
        page,
        width=300,
        height=40,
        values=[
            "Mouvetement",
            "Ckeck list",
            "Clok list"
        ],
        state="readonly"
    )
    select_type.set("Sélectionner un type")
    select_type.pack(anchor="w", padx=30)


    # ========================================================
    # AFFICHER LES FICHIERS
    # ========================================================

    def afficher_liste_fichiers():

        for widget in liste.winfo_children():
            widget.destroy()


        nombre = len(fichiers_selectionnes)
        if nombre == 0:
            nombre_fichiers.configure(
                text="0 fichier"
            )

            message = ctk.CTkLabel(
                liste,
                text="Aucun fichier sélectionné\n\nCliquez sur « Ajouter des fichiers » pour commencer.",
                text_color="#9CA3AF",
                font=ctk.CTkFont(
                    size=13
                )
            )

            message.pack(
                pady=80
            )
            return


        if nombre == 1:
            texte_nombre = "1 fichier"
        else:
            texte_nombre = f"{nombre} fichiers"


        nombre_fichiers.configure(
            text=texte_nombre
        )


        for index, chemin in enumerate(fichiers_selectionnes):
            nom = os.path.basename(chemin)
            ligne = ctk.CTkFrame(
                liste,
                fg_color="white",
                corner_radius=10,
                height=55
            )

            ligne.pack(
                fill="x",
                pady=4
            )
            ligne.pack_propagate(False)


            # Icône / extension
            extension = os.path.splitext(nom)[1].upper().replace(".", "")
            type_fichier = ctk.CTkLabel(
                ligne,
                text=extension,
                width=45,
                height=30,
                corner_radius=7,
                fg_color="#DCFCE7",
                text_color="#15803D",
                font=ctk.CTkFont(
                    size=11,
                    weight="bold"
                )
            )

            type_fichier.pack(
                side="left",
                padx=(10, 8)
            )


            # Nom du fichier
            label_nom = ctk.CTkLabel(
                ligne,
                text=nom,
                text_color="#374151",
                font=ctk.CTkFont(
                    size=13
                ),
                anchor="w"
            )

            label_nom.pack(
                side="left",
                fill="x",
                expand=True
            )


            # Supprimer
            def supprimer(index=index):

                fichiers_selectionnes.pop(index)
                afficher_liste_fichiers()


            bouton_supprimer = ctk.CTkButton(
                ligne,
                text="×",
                width=32,
                height=32,
                corner_radius=8,
                fg_color="#FEF2F2",
                hover_color="#FEE2E2",
                text_color="#DC2626",
                font=ctk.CTkFont(
                    size=18,
                    weight="bold"
                ),
                command=supprimer
            )

            bouton_supprimer.pack(
                side="right",
                padx=10
            )


    # ========================================================
    # BOUTON COMPILATION
    # ========================================================





    def lancer_la_compilation():
        type_compilation = select_type.get()

        if type_compilation == "Sélectionner un type":
            messagebox.showwarning(
                "Attention",
                "Veuillez sélectionner un type de compilation."
            )
            return

        bouton_compiler.configure(
            state="disabled",
            text="Compilation en cours..."
        )

        barre_progression.place(relx=0.98,y=20,anchor="ne")
        label_progression.place(relx=0.98,y=42,anchor="ne")

        
        barre_progression.set(0)


        label_progression.configure(text="Préparation de la compilation...")
        if not fichiers_selectionnes:
            messagebox.showwarning(
                "Aucun fichier",
                "Veuillez sélectionner au moins un fichier Excel."
            )
            return

        try:
            compilation = Workbook()

            # Supprimer la feuille vide créée automatiquement
            feuille_initiale = compilation.active
            compilation.remove(feuille_initiale)

            for index, chemin in enumerate(fichiers_selectionnes, start=1):

                fichier = os.path.basename(chemin)

                workbook_source = load_workbook(chemin)
                feuille_source = workbook_source.worksheets[0]

                nom_sheet = os.path.splitext(fichier)[0]
                nom_sheet = nom_sheet[:31]

                # Éviter deux feuilles avec le même nom
                nom_original = nom_sheet
                compteur = 1

                while nom_sheet in compilation.sheetnames:
                    suffixe = f"_{compteur}"
                    nom_sheet = nom_original[:31 - len(suffixe)] + suffixe
                    compteur += 1

                feuille_destination = compilation.create_sheet(nom_sheet)

                # Copier les cellules
                for ligne in feuille_source.iter_rows():

                    for cellule in ligne:
                        nouvelle_cellule = feuille_destination[cellule.coordinate]
                        nouvelle_cellule.value = cellule.value

                        if cellule.has_style:
                            nouvelle_cellule.font = copy(cellule.font)
                            nouvelle_cellule.fill = copy(cellule.fill)
                            nouvelle_cellule.border = copy(cellule.border)
                            nouvelle_cellule.alignment = copy(cellule.alignment)
                            nouvelle_cellule.number_format = cellule.number_format
                            nouvelle_cellule.protection = copy(cellule.protection)

                # Copier largeur des colonnes
                for colonne, dimension in feuille_source.column_dimensions.items():
                    feuille_destination.column_dimensions[colonne].width = dimension.width

                # Copier hauteur des lignes
                for ligne, dimension in feuille_source.row_dimensions.items():
                    feuille_destination.row_dimensions[ligne].height = dimension.height

                # Copier les cellules fusionnées
                for plage in feuille_source.merged_cells.ranges:
                    feuille_destination.merge_cells(str(plage))

                workbook_source.close()
                bouton_compiler.configure(
                    state="normal",
                    text="Lancer la compilation"
                )

                progression = index / len(fichiers_selectionnes)

                barre_progression.set(progression)
            
                pourcentage = int(progression * 100)

                label_progression.configure(text=f"{pourcentage} %")
                label_progression.configure(
                    text=f"Compilation du fichier {index} sur {len(fichiers_selectionnes)}..."
                )

                page.update()

       
                # Vider les fichiers affichés
                fichiers_selectionnes.clear()

                # Rafraîchir la liste
                afficher_liste_fichiers()
                

                barre_progression.place_forget()
                label_progression.place_forget()

                bouton_compiler.configure(
                    state="normal",
                    text="Lancer la compilation"
                )
            # Demander où enregistrer le résultat
            type_compilation = select_type.get()

            if type_compilation == "Mouvetement":
                nom_rapport = f"DAILY REPORT Mouvetement {date.today()}.xlsx"

            elif type_compilation == "Ckeck list":
                nom_rapport = f"DAILY REPORT Ckeck list {date.today()}.xlsx"

            elif type_compilation == "Clok list":
                nom_rapport = f"DAILY REPORT Clok list {date.today()}.xlsx"

            chemin_final = customtkinter.filedialog.asksaveasfilename(
                title="Enregistrer la compilation",
                defaultextension=".xlsx",
                filetypes=[
                    ("Fichier Excel", "*.xlsx")
                ],
                initialfile=nom_rapport

            )

            if not chemin_final:
                return

            compilation.save(chemin_final)

            messagebox.showinfo(
                "Compilation terminée",
                f"La compilation a été réalisée avec succès.\n\n"
                f"Fichiers compilés : {len(fichiers_selectionnes)}\n"
                f"Fichier créé :\n{chemin_final}"
            )

        except Exception as erreur:

            messagebox.showerror(
                "Erreur",
                f"Une erreur est survenue pendant la compilation :\n\n{erreur}"
            )




    bouton_compiler = ctk.CTkButton(
        page,
        text="Lancer la compilation",
        width=200,
        height=42,
        corner_radius=10,
        fg_color="#16A34A",
        hover_color="#15803D",
        command=lancer_la_compilation
    )

    bouton_compiler.pack(
        anchor="e",
        pady=(15, 0)
    )



    barre_progression = ctk.CTkProgressBar(
        page,
        width=500,
        height=12,
        corner_radius=6,
        progress_color="#16A34A"
    )

    label_progression = ctk.CTkLabel(
        page,
        text="0 %",
        font=("Arial", 12),
        text_color="#6B7280"
    )




    barre_progression.pack_forget()
    label_progression.pack_forget()




    # ========================================================
    # AFFICHAGE INITIAL
    # ========================================================
    afficher_liste_fichiers()

















def afficher_historyque():
    # Nettoyer le conteneur principal
    for widget in conteneur1.winfo_children():
        widget.destroy()

    













# LANCEMENT
# ============================================================
fenetre.mainloop()
